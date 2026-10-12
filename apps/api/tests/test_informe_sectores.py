"""T5 sector-grouped report tests (TDD).

When opportunities carry a ``sector``, ``generate_report_html`` groups cards
by sector in ``SECTOR_TAXONOMY_V1`` order (``otros`` last) with a per-group
header (count badge + secondary meta row with top amounts/dates when
available). With no sector on any row the output keeps STRUCTURAL
compatibility with the legacy flat report (same data, same order, no
sector-group markup) — not byte-identity: the visual polish pass
(brand-bar air, sector band, equal-height cards with clamped summaries,
print-friendly 1px borders) intentionally changed the CSS/HTML shell, so
the old byte-identical pin was replaced by this structural contract.
``export_xlsx`` keeps the default sheet plus a ``sector`` column and
``export_pdf`` still derives from the same HTML pipeline.

Zero network: fixtures use ``application_url=None`` so the sync URL
reachability check short-circuits before any I/O.
"""

from __future__ import annotations

from datetime import datetime
from types import SimpleNamespace

from app.services.export import export_pdf, export_xlsx, generate_report_html


def _opp(title, sector=None, **overrides):
    base = dict(
        source_id="src-1",
        status="open",
        country="Colombia",
        title=title,
        entity="Test Entity",
        summary="Summary for %s" % title,
        description="",
        close_date=None,
        funding_amount_raw="",
        funding_amount_value=None,
        categories=[],
        official_url="",
        application_url=None,
    )
    base.update(overrides)
    opp = SimpleNamespace(**base)
    # Attach ``sector`` only when explicitly given so compat fixtures can
    # omit the attribute entirely (like pre-migration rows).
    if sector is not None or "sector" in overrides:
        opp.sector = sector
    return opp


def _sectorized():
    return [
        _opp("Salud rural grant", sector="salud",
             funding_amount_raw="$50.000.000",
             close_date=datetime(2026, 11, 30)),
        _opp("Beca docente", sector="educación",
             funding_amount_raw="$10.000.000",
             close_date=datetime(2026, 12, 15)),
        _opp("Hospital equipment fund", sector="salud",
             funding_amount_raw="$80.000.000",
             close_date=datetime(2026, 10, 20)),
        _opp("Unclassified call"),
    ]


class TestSectorGrouping:
    def test_groups_by_sector_in_taxonomy_order(self):
        html = generate_report_html("Report", SimpleNamespace(name="Org"), _sectorized())
        edu = html.find("educaci")
        health = html.find("salud")
        assert edu >= 0 and health >= 0
        # SECTOR_TAXONOMY_V1 ranks educacion before salud.
        assert edu < health
        for title in ("Salud rural grant", "Beca docente", "Hospital equipment fund"):
            assert title in html

    def test_group_header_shows_count_amounts_and_dates(self):
        html = generate_report_html("Report", SimpleNamespace(name="Org"), _sectorized())
        assert "2" in html  # salud group count surfaces in its header
        assert "$50.000.000" in html or "$80.000.000" in html
        assert "2026-10-20" in html or "2026-11-30" in html

    def test_otros_group_sorts_last(self):
        opps = [
            _opp("Misc call", sector="otros"),
            _opp("Andes trek fund", sector="turismo"),
        ]
        html = generate_report_html("Report", SimpleNamespace(name="Org"), opps)
        turismo = html.find("turismo")
        otros = html.find("otros")
        assert turismo >= 0 and otros >= 0
        assert turismo < otros

    def test_unsectorized_row_falls_into_otros_group(self):
        html = generate_report_html("Report", SimpleNamespace(name="Org"), _sectorized())
        assert "Unclassified call" in html
        assert "otros" in html


class TestSectorCompat:
    def test_no_sectors_output_has_no_grouping_markup(self):
        opps = [_opp("Call A"), _opp("Call B")]
        html = generate_report_html("Report", SimpleNamespace(name="Org"), opps)
        assert "sector-group" not in html
        assert "Todas las convocatorias" in html
        assert "Call A" in html and "Call B" in html

    def test_no_sectors_keeps_data_and_order(self):
        """Structural compat: same cards, same input order, flat grid."""
        opps = [_opp("Call A"), _opp("Call B")]
        html = generate_report_html("Report", SimpleNamespace(name="Org"), opps)
        assert html.find("Call A") < html.find("Call B")
        assert "story-grid" in html
        assert "data-sector" not in html


class TestReportPolishStructure:
    """Layout contract for the visual polish pass (structure, not pixels)."""

    def _html(self):
        return generate_report_html("Report", SimpleNamespace(name="Org"), _sectorized())

    def test_brand_bar_has_air_and_hierarchy(self):
        html = self._html()
        assert "brand-bar" in html and "brand-name" in html and "brand-tagline" in html
        assert "gap:" in html or "gap :" in html  # breathing room by the shield lockup

    def test_sector_header_has_count_badge_and_meta_row(self):
        html = self._html()
        assert "sector-group__count" in html
        assert "2 oportunidades" in html  # salud group count badge
        assert "sector-group__meta" in html

    def test_cards_pin_footer_and_clamp_summary(self):
        html = self._html()
        assert "story-card__foot" in html
        assert "margin-top: auto" in html or "margin-top:auto" in html  # CTA aligned to foot
        assert "line-clamp" in html  # long summaries never blow up the card

    def test_meta_grid_and_badges_are_aligned(self):
        html = self._html()
        assert "story-card__meta-grid" in html
        for label in ("Entidad", "Cierre", "Monto"):
            assert label in html
        assert "badge--open" in html

    def test_print_is_flat_and_unbroken(self):
        html = self._html()
        assert "@media print" in html
        assert "break-inside: avoid" in html


class TestSectorXlsx:
    def test_xlsx_keeps_default_sheet_and_adds_sector_column(self):
        from openpyxl import load_workbook
        import io

        data = export_xlsx(_sectorized())
        workbook = load_workbook(io.BytesIO(data))
        sheet = workbook["Convocatorias"]
        header = [cell.value for cell in sheet[1]]
        assert "Sector" in header
        sector_idx = header.index("Sector")
        values = [row[sector_idx].value for row in sheet.iter_rows(min_row=2)]
        assert values == ["salud", "educación", "salud", None]


class TestSectorPdf:
    def test_pdf_is_non_empty(self):
        pdf = export_pdf("Report", SimpleNamespace(name="Org"), _sectorized())
        assert isinstance(pdf, bytes) and len(pdf) > 0
